from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("TraceMate_Technical_Overview.pptx")

NAVY = RGBColor(12, 23, 38)
NAVY_LIGHT = RGBColor(23, 40, 59)
BLUE = RGBColor(42, 131, 194)
TEAL = RGBColor(0, 164, 153)
WHITE = RGBColor(247, 249, 252)
MUTED = RGBColor(171, 184, 198)
GRAY = RGBColor(92, 109, 125)
PANEL = RGBColor(20, 35, 52)
LINE = RGBColor(92, 120, 146)
GREEN = RGBColor(94, 190, 154)


def color(value):
    return RGBColor(value[0], value[1], value[2]) if isinstance(value, tuple) else value


def set_background(slide):
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY


def add_text(slide, text, x, y, w, h, size, fill=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = fill
    return box


def add_title(slide, number, title, subtitle=None):
    add_text(slide, f"0{number}", 0.62, 0.39, 0.55, 0.28, 10, TEAL, True)
    add_text(slide, title, 0.62, 0.67, 11.8, 0.55, 26, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, 0.64, 1.23, 11.6, 0.34, 10.5, MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.62), Inches(1.67), Inches(12.05), Inches(0.025))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()


def add_footer(slide):
    add_text(slide, "TRACEMATE  /  TECHNICAL OVERVIEW", 0.62, 7.14, 4.5, 0.16, 7.5, MUTED, True)
    add_text(slide, "Shareable overview", 10.75, 7.14, 1.9, 0.16, 7.5, MUTED, False, PP_ALIGN.RIGHT)


def add_panel(slide, x, y, w, h, title, body, accent=BLUE):
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = LINE
    panel.line.width = Pt(0.65)
    marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    marker.fill.solid()
    marker.fill.fore_color.rgb = accent
    marker.line.fill.background()
    add_text(slide, title, x + 0.25, y + 0.22, w - 0.45, 0.28, 13, WHITE, True)
    add_text(slide, body, x + 0.25, y + 0.56, w - 0.45, h - 0.72, 10.5, MUTED)


def add_node(slide, x, y, w, h, title, details, accent):
    node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    node.fill.solid()
    node.fill.fore_color.rgb = NAVY_LIGHT
    node.line.color.rgb = accent
    node.line.width = Pt(1.35)
    add_text(slide, title, x + 0.17, y + 0.16, w - 0.34, 0.28, 12, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, details, x + 0.14, y + 0.52, w - 0.28, h - 0.65, 8.8, MUTED, False, PP_ALIGN.CENTER)


def add_connector(slide, x1, y1, x2, y2, label, label_x, label_y, label_w):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = LINE
    connector.line.width = Pt(1.25)
    connector.line.end_arrowhead = True
    add_text(slide, label, label_x, label_y, label_w, 0.24, 8.3, MUTED, False, PP_ALIGN.CENTER)


def add_bullet_list(slide, items, x, y, w, font_size=13):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(3.8))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = WHITE
        paragraph.space_after = Pt(15)
        paragraph.bullet = True
    return box


def build():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # Slide 1: title
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.62), Inches(0.62), Inches(0.12), Inches(1.0))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_text(slide, "TraceMate", 0.98, 0.66, 5.5, 0.63, 30, WHITE, True)
    add_text(slide, "Guided diagnostics and trace export for vehicle testing", 1.0, 1.38, 7.5, 0.42, 15, MUTED)
    add_text(slide, "A focused Android workflow that connects the technician to the head unit and a headless CAPTURE_TOOL_PLACEHOLDER Mac.", 1.0, 2.23, 8.9, 0.54, 18, WHITE)
    add_panel(slide, 1.0, 3.35, 3.5, 1.6, "AUTOMATE", "Connection readiness, firewall preparation, ADB validation, and repeatable CAPTURE_TOOL_PLACEHOLDER control.", BLUE)
    add_panel(slide, 4.9, 3.35, 3.5, 1.6, "SIMPLIFY", "Select and export head-unit traces to a USB stick through one visible, guided flow.", TEAL)
    add_panel(slide, 8.8, 3.35, 3.5, 1.6, "PROTECT", "Explicit user actions, observable progress, retries, and recovery for long-running transfers.", GREEN)
    add_text(slide, "PROJECT OVERVIEW", 1.0, 6.45, 2.4, 0.2, 8.5, TEAL, True)
    add_text(slide, "Android app  |  Head unit diagnostics  |  CAPTURE_TOOL_PLACEHOLDER capture bridge", 1.0, 6.72, 6.5, 0.24, 10, MUTED)
    add_footer(slide)

    # Slide 2: architecture
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 2, "System topology", "Two independent network paths keep vehicle diagnostics and CAPTURE_TOOL_PLACEHOLDER control deliberate and traceable.")
    add_node(slide, 0.75, 2.18, 2.55, 1.35, "HEAD UNIT", "SSH + ADB\nWLAN hotspot\nUSB host", BLUE)
    add_node(slide, 5.35, 3.1, 2.62, 1.42, "ANDROID APP", "TraceMate\nWorkflow controller\nVisible status", TEAL)
    add_node(slide, 10.0, 1.95, 2.55, 1.42, "MAC MINI", "Headless CAPTURE_TOOL_PLACEHOLDER agent\nSSH / SCP endpoint\nCAPTURE_TOOL_PLACEHOLDER captures", GREEN)
    add_node(slide, 10.0, 4.55, 2.55, 1.1, "MOBILE_DEVICE_PLACEHOLDER", "USB-connected to Mac\nCAPTURE_SESSION_PLACEHOLDER target", RGBColor(214, 166, 71))
    add_node(slide, 0.75, 5.1, 2.55, 1.05, "USB STICK", "Physically connected\nto head unit", RGBColor(214, 166, 71))
    add_connector(slide, 3.3, 2.85, 5.35, 3.64, "WLAN hotspot  |  SSH + ADB", 3.45, 2.65, 1.85)
    add_connector(slide, 7.97, 3.55, 10.0, 2.82, "Android Ethernet tethering  |  SSH + SCP", 8.0, 2.46, 2.15)
    add_connector(slide, 10.95, 3.37, 11.1, 4.55, "USB", 11.25, 3.83, 0.6)
    add_connector(slide, 2.0, 3.53, 2.0, 5.1, "Local copy", 2.17, 4.15, 0.9)
    add_text(slide, "Store-and-forward: the Mac has no direct head-unit route. Android retrieves completed CAPTURE_TOOL_PLACEHOLDER captures, then uploads them to the head-unit USB.", 0.78, 6.5, 11.75, 0.31, 10.3, MUTED)
    add_footer(slide)

    # Slide 3: workflows
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 3, "Automated testing and export workflows", "The app exposes the current readiness and next safe action instead of requiring technicians to coordinate individual tools.")
    add_panel(slide, 0.72, 2.02, 3.8, 3.85, "1  HEAD-UNIT READINESS", "Connect to the head-unit WLAN via QR-guided setup.\n\nAutomatically validate:\nSSH -> scoped firewall -> ADB TCP -> ADB handshake -> shell access.\n\nDisplay E-Release only when the connection path is ready.", BLUE)
    add_panel(slide, 4.77, 2.02, 3.8, 3.85, "2  CAPTURE_TOOL_PLACEHOLDER CAPTURE", "Reach the headless Mac over Android Ethernet tethering.\n\nGuide required manual steps: MobileDevicePlaceholder trust, Bluetooth state, and CAPTURE_SESSION_PLACEHOLDER confirmation.\n\nThe Mac agent owns the CAPTURE_TOOL_PLACEHOLDER process and persists capture-job state.", TEAL)
    add_panel(slide, 8.82, 2.02, 3.8, 3.85, "3  USB EXPORT", "Discover traces and writable USB storage on the head unit.\n\nSelect archives, show progress, allow cancellation, and report success, partial failure, or recovery.\n\nExport CAPTURE_TOOL_PLACEHOLDER captures through Android store-and-forward when required.", GREEN)
    add_text(slide, "Technician value: fewer tool switches, fewer manual command sequences, and a repeatable handover of diagnostic artifacts.", 0.75, 6.32, 11.7, 0.35, 12.5, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # Slide 4: reliability
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 4, "Operational reliability by design", "Automation is constrained to known workflows, with visible state and recovery for field use.")
    add_bullet_list(slide, [
        "Scoped firewall rules limit ADB access to the head-unit WLAN client range.",
        "All state-changing actions require an explicit technician trigger and show progress or failure feedback.",
        "USB exports use staged partial files and atomic rename to avoid presenting incomplete archives as final output.",
        "Long-running exports continue in a foreground service and can reattach to a remote worker after app restart.",
        "The headless Mac agent allows one managed CAPTURE_TOOL_PLACEHOLDER job, preserves job state, and keeps its RPC interface loopback-only.",
    ], 0.82, 2.08, 7.0, 13)
    add_panel(slide, 8.35, 2.04, 3.95, 1.15, "PRIMARY OUTCOME", "Reduce time from vehicle connection to usable trace package.", TEAL)
    add_panel(slide, 8.35, 3.52, 3.95, 1.15, "SECONDARY OUTCOME", "Make failures actionable before a technician leaves the vehicle.", BLUE)
    add_panel(slide, 8.35, 5.0, 3.95, 1.15, "CURRENT SCOPE", "Guided workflow and diagnostics; not yet a one-click production process.", GREEN)
    add_text(slide, "TraceMate makes complex, multi-device testing easier to operate without hiding the safety-critical steps.", 0.82, 6.5, 11.55, 0.33, 11.5, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide)

    presentation.core_properties.title = "TraceMate Technical Overview"
    presentation.core_properties.subject = "Android, head unit, Mac mini, CAPTURE_TOOL_PLACEHOLDER, and USB trace export architecture"
    presentation.core_properties.author = "TraceMate"
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
